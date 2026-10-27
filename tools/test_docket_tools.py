"""Regression tests for the Docket tools: capture splitting, extraction, the question floor,
export, hooks, scope.

Usage: python tools/test_docket_tools.py
Builds a throwaway bundle (PDF, HTML, TXT, DOCX and an unsupported file) and a small task-page
capture in a temp directory; nothing under tasks/ or dist/ is touched. Run it after changing any
tool.
"""
from __future__ import annotations

import contextlib
import io
import json
import re
import shutil
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

TOOLS = Path(__file__).resolve().parent
sys.path.insert(0, str(TOOLS))

import agent_scope  # noqa: E402
import check_questions as cq  # noqa: E402
import extract_text  # noqa: E402
import form_numbers as fn  # noqa: E402
import split_capture  # noqa: E402

FILLER = ("The company operates a regional freight network across four states and reports results "
          "each quarter. Management discussed fuel costs, driver retention and terminal utilisation. ")
GOOD = """---
id: q01
status: draft
origin: user
asks: 2
answer_type: yes-no; number
documents: D01, D02
tags: Cross-document join; Temporal alignment
levers: Scatter the pieces; Calculate across documents
revision: 0
---

## Question

Did the freight company's actual revenue for the second calendar quarter of 2025 miss or beat
the revenue it had forecast for that quarter, and by how many million dollars, to one
decimal place?

## Answer

Missed, by $2.8 million

## Components

- c1: Missed
- c2: $2.8 million

## Reasoning

Step 1. Source fact (Doc 2, first-quarter release): The forecast of $412.6 million is the revenue
the company expected for the second quarter.

Step 2. Source fact (Doc 1, revenue table): the actual revenue of 409.8 million dollars for Q2 2025 is in
the second-quarter table.

Step 3. Calculation (Doc 1 and Doc 2): actual minus forecast, in millions of dollars: 409.8 - 412.6 = -2.8, so revenue missed the forecast.

## Evidence

- D01 (p. 1) [c1, c2]: "expects second-quarter revenue of $412.6 million"
- D02 (revenue table) [c1, c2]: "Q2 2025 | 409.8"

## Targeted failure modes

### Failure mode 1
Trap: Figure that must be computed
How a careless solver fails: Misapplies a rule or premise

Why these are traps: The second-quarter release (Doc 1) calls revenue "in line" with the forecast,
so a careless solver answers that it met the forecast. The forecast of 412.6 sits only in the
first-quarter release; against the table's 409.8, revenue missed by 2.8 million dollars.

## Lever checks

- One document is not enough: the forecast and the actual are in different releases.

## Blind solve

## Review

## Feedback log
"""


def make_bundle(root: Path) -> Path:
    import pymupdf
    from docx import Document
    bundle = root / "tasks" / "freight"
    docs = bundle / "docs"
    docs.mkdir(parents=True)
    pdf = pymupdf.open()
    page = pdf.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), FILLER * 2
                        + "Harlow Freight expects second-quarter revenue of $412.6 million, which it "
                          "described as in line with plans.", fontsize=10)
    pdf.new_page().insert_text((50, 60), "Appendix", fontsize=10)
    pdf.save(docs / "a-release-q1.pdf")
    (docs / "b-release-q2.html").write_bytes(
        b"<html><body><p>Revenue for the second quarter was <b>in line</b> with our forecast.</p>"
        b"<table><tr><th>Quarter</th><th>Revenue ($ millions)</th></tr>"
        b"<tr><td>Q2 2025</td><td>409.8</td></tr></table><script>var x=1;</script></body></html>")
    (docs / "c-notes.txt").write_bytes("Analyst note\r\nOperating margin was 14.6% in 2025.\r\n".encode())
    memo = Document()
    memo.add_paragraph("Board memo on terminal leases.")
    table = memo.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Terminal", "Lease ends 2027"
    memo.save(docs / "d-memo.docx")
    (docs / "e-sheet.xlsx").write_bytes(b"PK\x03\x04 not really a workbook")
    (docs / ".hidden").write_bytes(b"ignored")
    return bundle


class DocketTools(unittest.TestCase):
    def setUp(self):
        self.root = Path(tempfile.mkdtemp(prefix="docket-test-"))
        self.bundle = make_bundle(self.root)
        self.report = extract_text.run(self.bundle)
        self.manifest = json.loads((self.bundle / "text" / "manifest.json").read_bytes())
        (self.bundle / "questions").mkdir()

    def tearDown(self):
        shutil.rmtree(self.root, ignore_errors=True)

    def write(self, text: str, name: str = "q01") -> Path:
        path = self.bundle / "questions" / f"{name}.md"
        path.write_bytes(text.encode("utf-8"))
        return path

    def codes(self, text: str, level: str = "FAIL") -> set[str]:
        _, _, found = cq.check(self.write(text), cq.Bundle(self.bundle))
        return {code for lvl, code, _ in found if lvl == level}

    def write_platform_numbers(self):
        # Intentionally different from local Dnn order.
        (self.bundle / "notes").mkdir(exist_ok=True)
        (self.bundle / "notes" / "intake.json").write_bytes(json.dumps({"schema": 1, "documents": [
            {"number": 1, "source": "docs/b-release-q2.html"},
            {"number": 2, "source": "docs/a-release-q1.pdf"},
            {"number": 3, "source": "docs/c-notes.txt"}]}).encode("utf-8"))

    # extract_text.py
    def test_ids_follow_sorted_names_and_formats(self):
        sources = {k: v["source"] for k, v in self.manifest.items()}
        self.assertEqual(sources, {"D01": "docs/a-release-q1.pdf", "D02": "docs/b-release-q2.html",
                                   "D03": "docs/c-notes.txt", "D04": "docs/d-memo.docx",
                                   "D05": "docs/e-sheet.xlsx"})
        self.assertEqual(self.manifest["D05"]["method"], "unsupported")
        self.assertFalse((self.bundle / "text" / "D05.txt").exists())

    def test_pdf_pages_thin_pages_and_layout(self):
        text = (self.bundle / "text" / "D01.txt").read_bytes().decode("utf-8")
        self.assertIn("=== page 1 ===", text)
        self.assertIn("=== page 2 ===", text)
        self.assertNotIn("\r", text)
        self.assertEqual(self.manifest["D01"]["thin_pages"], [2])
        if shutil.which("pdftotext"):
            self.assertTrue((self.bundle / "text" / "D01.layout.txt").exists())

    def test_html_rows_docx_tables_and_text(self):
        html = (self.bundle / "text" / "D02.txt").read_bytes().decode("utf-8")
        self.assertIn("Q2 2025 | 409.8", html)
        self.assertIn("Revenue for the second quarter was in line with our forecast.", html)
        self.assertNotIn("var x", html)
        self.assertIn("Terminal | Lease ends 2027", (self.bundle / "text" / "D04.txt").read_bytes().decode())
        self.assertNotIn(b"\r", (self.bundle / "text" / "D03.txt").read_bytes())

    def test_rerun_keeps_ids(self):
        self.assertTrue(all("[unchanged]" in line for line in extract_text.run(self.bundle)))
        (self.bundle / "docs" / "0-new.txt").write_bytes(b"A late addition.")
        (self.bundle / "docs" / "c-notes.txt").write_bytes(b"Operating margin was 14.9% in 2025.")
        (self.bundle / "docs" / "d-memo.docx").unlink()
        extract_text.run(self.bundle)
        manifest = json.loads((self.bundle / "text" / "manifest.json").read_bytes())
        self.assertEqual(manifest["D06"]["source"], "docs/0-new.txt")
        self.assertEqual(manifest["D03"]["source"], "docs/c-notes.txt")
        self.assertIn("14.9%", (self.bundle / "text" / "D03.txt").read_bytes().decode())
        self.assertTrue(manifest["D04"]["missing"])

    # check_questions.py
    def test_good_question_has_no_fail(self):
        self.assertEqual(self.codes(GOOD), set())

    def test_dot_leader_quote_is_one_passage(self):
        # Federal Register timetables print dot leaders; a quote that holds them verbatim is
        # one contiguous passage, while dots that do not occur in the document stay a join.
        (self.bundle / "docs" / "c-notes.txt").write_bytes(
            b"Timetable: NPRM Comment Period End............. 02/05/24 End of Extended Comment "
            b"Period...... 04/05/24 Final Rule 01/00/25\r\n")
        extract_text.run(self.bundle)
        leader = GOOD.replace('"Q2 2025 | 409.8"\n',
                              '"Q2 2025 | 409.8"\n- D03 (timetable) [c1]: '
                              '"End of Extended Comment Period...... 04/05/24"\n')
        self.assertNotIn("evidence", self.codes(leader))
        joined = leader.replace('Period...... 04/05/24"', 'Period... 04/05/24"')
        self.assertIn("evidence", self.codes(joined))

    def test_pipe_table_row_matches_platform_text(self):
        # A complete capture renders hidden tables as Markdown pipe rows; the platform's own text
        # shows the cells space-separated, so either form of the quote must match the row.
        (self.bundle / "docs" / "c-notes.txt").write_bytes(
            b"|  |  |  |\n| --- | --- | --- |\n| Dividends Paid | (538) | (528) |\n")
        extract_text.run(self.bundle)
        row = GOOD.replace('"Q2 2025 | 409.8"\n',
                           '"Q2 2025 | 409.8"\n- D03 (cash flow) [c1]: "Dividends Paid (538) (528)"\n')
        self.assertNotIn("evidence", self.codes(row))
        piped = row.replace('"Dividends Paid (538) (528)"', '"Dividends Paid | (538) | (528)"')
        self.assertNotIn("evidence", self.codes(piped))
        wrong = row.replace('"Dividends Paid (538) (528)"', '"Dividends Paid (538) (529)"')
        self.assertIn("evidence", self.codes(wrong))

    def test_native_pipe_row_needs_its_pipes(self):
        # A document whose own text is a Markdown table ("|---|" rules, no spaces) keeps its pipes on the
        # page, and the form rejected every pipe-free quote of such rows (aq-06840 q01 and q02 rev 0,
        # user screenshots 2026-10-18). The quote must copy the row as printed; the spaced rows the older
        # captures drew themselves (test above) keep the pipe-free allowance.
        (self.bundle / "docs" / "c-notes.txt").write_bytes(
            b"|Item|2025|2024|\n|---|---|---|\n|Dividends Paid|(538)|(528)|\n")
        extract_text.run(self.bundle)
        piped = GOOD.replace('"Q2 2025 | 409.8"\n',
                             '"Q2 2025 | 409.8"\n- D03 (cash flow) [c1]: "|Dividends Paid|(538)|(528)|"\n')
        self.assertNotIn("evidence", self.codes(piped))
        bare = piped.replace('"|Dividends Paid|(538)|(528)|"', '"Dividends Paid (538) (528)"')
        self.assertIn("evidence", self.codes(bare))

    def test_floor_catches_each_rule(self):
        cases = {
            "evidence": GOOD.replace("$412.6 million\"", "$412.9 million\""),
            "documents": GOOD.replace("documents: D01, D02", "documents: D01"),
            "self-contained": GOOD.replace("to one\ndecimal place?", "to one\ndecimal place, as on page 3?"),
            "essay": GOOD.replace("Did the freight", "Explain why the freight"),
            "giveaway": GOOD.replace("to one\ndecimal place?", "to one\ndecimal place (missed, by $2.8 million)?"),
            "decoy": GOOD.replace("Did the freight company's", "Which quarter shows the freight company's"),
            "tags": GOOD.replace("tags: Cross-document join; Temporal alignment", "tags: Cross-document join"),
            "status": GOOD.replace("status: draft", "status: ready"),
            "question-length": GOOD.replace("to one\ndecimal place?", "to one decimal place? " + "word " * 60),
            "fields": GOOD.replace("Missed, by $2.8 million", "TODO"),
            "asks": GOOD.replace("asks: 2", "asks: 3"),
        }
        for code, text in cases.items():
            with self.subTest(code=code):
                self.assertIn(code, self.codes(text))
        # The task form's quality checks (captured 2026-10-04).
        question = ("Did the freight company's actual revenue for the second calendar quarter of 2025 miss or beat\n"
                    "the revenue it had forecast for that quarter, and by how many million dollars, to one\n"
                    "decimal place?")
        self.assertIn(question, GOOD)
        form = {
            "semicolon in the answer": ("short-answer", GOOD.replace("Missed, by $2.8 million", "Missed; by $2.8 million")),
            "answer over 10 words": ("short-answer", GOOD.replace(
                "Missed, by $2.8 million", "Missed, by $2.8 million, less than the company had planned for")),
            "question under 8 words": ("question-length", GOOD.replace(question, "Did revenue miss the forecast?")),
            "answer part in the question": ("giveaway", GOOD.replace("miss or beat", "missed or beat")),
            "quote under 12 characters": ("evidence", GOOD.replace('"Q2 2025 | 409.8"', '"409.8"')),
            "quote joined by an ellipsis": ("evidence", GOOD.replace(
                '"expects second-quarter revenue of $412.6 million"', '"expects second-quarter ... $412.6 million"')),
            "no components": ("components", GOOD.replace("- c1: Missed\n- c2: $2.8 million\n", "")),
            "component not in the answer": ("components", GOOD.replace("- c1: Missed", "- c1: Beat")),
            "number in no component": ("components", GOOD.replace("- c2: $2.8 million", "- c2: million")),
            "quote without a component": ("components", GOOD.replace(' [c1, c2]: "Q2', ': "Q2')),
            "component without a quote": ("components", GOOD.replace("[c1, c2]", "[c2]")),
            "fewer components than asks": ("components", GOOD.replace("- c2: $2.8 million\n", "").replace(", c2]", "]")),
            "parts not told apart": ("components", GOOD.replace("Missed, by $2.8 million", "Missed $2.8 million")),
        }
        for name, (code, text) in form.items():
            with self.subTest(form=name):
                self.assertIn(code, self.codes(text))
        # One bare count may go uncited.
        count = GOOD.replace("- c2: $2.8 million", "- c2: $2.8 million\n- c3: 2 quarters").replace(
            "Missed, by $2.8 million", "Missed, by $2.8 million, 2 quarters")
        self.assertNotIn("components", self.codes(count))

    def test_floor_warnings(self):
        self.write_platform_numbers()  # the form's reasoning number check needs the page's Doc numbers
        copied = GOOD.replace("Did the freight", "Harlow Freight expects second-quarter revenue of $412.6 million; did the")
        self.assertIn("copied-phrasing", self.codes(copied, "WARN"))
        rate = GOOD.replace("Did the freight company's actual revenue", "How much did the operating margin change versus revenue")
        self.assertIn("points-or-percent", self.codes(rate, "WARN"))
        loose = GOOD.replace(", to one\ndecimal place?", "?").replace("calendar ", "")
        self.assertTrue({"precision", "period"} <= self.codes(loose, "WARN"))
        self.assertIn("units", self.codes(GOOD.replace("Missed, by $2.8 million", "2.8"), "WARN"))
        self.assertIn("documents", self.codes(GOOD.replace("documents: D01, D02", "documents: D01, D02, D03, D04, D05"), "WARN"))
        self.assertIn("origin", self.codes(GOOD.replace("origin: user", "origin: claude"), "WARN"))
        self.assertIn("guessable", self.codes(GOOD.replace("answer_type: yes-no; number", "answer_type: yes-no"), "WARN"))
        # The task form's advisory checks (captured 2026-10-04).
        listed = GOOD.replace("answer_type: yes-no; number", "answer_type: list")
        advisory = {
            "question-length": GOOD.replace("to one\ndecimal place?", "to one decimal place? " + "word " * 30),
            "style": GOOD.replace("= -2.8, so", "= -2.8 — so"),
            "format": listed,
            "list-contract": listed.replace("Missed, by $2.8 million", "Alpha, Beta, Gamma, Delta"),
            "partial": GOOD.replace("miss or beat", "miss or beat at least one"),
            "world": GOOD.replace("Did the freight", "In this bundle, did the freight"),
        }
        for code, text in advisory.items():
            with self.subTest(advisory=code):
                self.assertIn(code, self.codes(text, "WARN"))
                self.assertNotIn(code, self.codes(text))
        self.assertIn("style", self.codes(GOOD.replace("so revenue missed", "it is worth noting that revenue missed"), "WARN"))
        ordered = listed.replace("to one\ndecimal place?", "to one\ndecimal place, if any, in any order?")
        self.assertNotIn("format", self.codes(ordered, "WARN"))
        # aq-03736 q02 (2026-10-13): the form's rule-10 advisory called "in any order" an answer-protocol tail.
        self.assertIn("world", self.codes(ordered, "WARN"))
        sorted_list = listed.replace("to one\ndecimal place?", "to one\ndecimal place, if any? List them alphabetically.")
        self.assertNotIn("format", self.codes(sorted_list, "WARN"))
        self.assertNotIn("world", self.codes(sorted_list, "WARN"))
        # Measured on aq-03170 q01 (2026-10-12): an order the form did not read as one, a cue clause that
        # measured a tier easier than the structural wording, and an ordinal the form took for a pointer.
        bare = listed.replace("to one\ndecimal place?", "to one\ndecimal place, if any? Earliest quarter first.")
        self.assertIn("format", self.codes(bare, "WARN"))
        cued = GOOD.replace("to one\ndecimal place?", "to one\ndecimal place, including restated figures?")
        self.assertIn("cue", self.codes(cued, "WARN"))
        # Measured on aq-03169 q02 (2026-10-13): a completeness word is a cue too, and a list gold written as
        # labelled sentences failed the form's blocking "Short answer" check.
        searched = GOOD.replace("Did the freight", "Anywhere in its filings, did the freight")
        self.assertIn("cue", self.codes(searched, "WARN"))
        labelled = listed.replace("Missed, by $2.8 million", "Result: missed. Size: $2.8 million")
        self.assertIn("short-answer", self.codes(labelled, "WARN"))
        self.assertNotIn("short-answer", self.codes(listed, "WARN"))
        pointed = GOOD.replace("Did the freight", "Going by the second table, did the freight")
        self.assertIn("pointer", self.codes(pointed, "WARN"))
        self.assertEqual(self.codes(GOOD, "WARN"), set())

    def test_floor_reads_the_2026_10_16_guide(self):
        # The guide's newer levers, its "two together land Medium or Hard", and send-back reasons it lists (sections
        # 5 and 7, saved in .claude/memory/reference/platform-guide-2026-10-16.md).
        self.write_platform_numbers()
        newer = GOOD.replace("levers: Scatter the pieces; Calculate across documents",
                             "levers: Restated figures; Ask for a precision the document doesn’t print")
        self.assertNotIn("levers", self.codes(newer, "WARN"))
        self.assertIn("levers", self.codes(GOOD.replace("Calculate across documents", "Hidden trap"), "WARN"))
        hard = GOOD.replace("revision: 0", "revision: 0\ndifficulty_target: hard")
        self.assertNotIn("levers", self.codes(hard, "WARN"))
        self.assertIn("levers", self.codes(hard.replace("Scatter the pieces; ", ""), "WARN"))
        cases = {
            "hint": GOOD.replace("actual revenue", "actual revenue, as originally reported,"),
            "gold-extra": GOOD.replace("Missed, by $2.8 million", "Missed, by about $2.8 million"),
            "selector": GOOD.replace("for the second calendar quarter of 2025",
                                     "for the calendar quarter of 2025 in which its depot opened"),
            "literal": GOOD.replace("the revenue it had forecast", "the revenue its latest release had forecast"),
        }
        for code, text in cases.items():
            with self.subTest(code=code):
                self.assertIn(code, self.codes(text, "WARN"))
                self.assertNotIn(code, self.codes(text))
        bracketed = GOOD.replace("Missed, by $2.8 million", "Missed, by $2.8 million (Q2)")
        self.assertIn("gold-extra", self.codes(bracketed, "WARN"))
        asked = bracketed.replace("to one\ndecimal place?", "to one\ndecimal place? Give the quarter in brackets.")
        self.assertNotIn("gold-extra", self.codes(asked, "WARN"))

    def test_list_limits_count_answer_items_not_evidence_components(self):
        def listed(answer, comp_text):
            return GOOD.replace("answer_type: yes-no; number", "answer_type: list").replace(
                "Missed, by $2.8 million", answer).replace(
                "- c1: Missed\n- c2: $2.8 million", comp_text)

        thirteen = ", ".join(f"Item{n}" for n in range(1, 14))
        self.assertIn("short-answer", self.codes(listed(thirteen, f"- c1: {thirteen}")))
        twelve = ", ".join(f"Item{n}" for n in range(1, 13))
        self.assertNotIn("short-answer", self.codes(listed(twelve, f"- c1: {twelve}")))
        four = "Alpha, Beta, Gamma, Delta"
        compressed = listed(four, f"- c1: {four}")
        self.assertIn("list-contract", self.codes(compressed, "WARN"))
        partial = compressed.replace("miss or beat", "miss or beat at least two")
        self.assertIn("partial", self.codes(partial, "WARN"))
        # Splitting one item's evidence into several components cannot inflate its count.
        single = listed('"Alpha Beta Gamma Delta"', "- c1: Alpha\n- c2: Beta\n- c3: Gamma\n- c4: Delta")
        self.assertNotIn("list-contract", self.codes(single, "WARN"))
        self.assertNotIn("partial", self.codes(single.replace("miss or beat", "miss or beat at least one"), "WARN"))
        long_item = " ".join("abcdefghijk")
        self.assertIn("answer-length", self.codes(listed(long_item, "- c1: a\n- c2: b"), "WARN"))

    def test_list_boundaries_preserve_numbers_dates_and_entity_names(self):
        cases = {
            "1,000, 2,000": ["1,000", "2,000"],
            "October 5, 2026, October 6, 2026": ["October 5, 2026", "October 6, 2026"],
            "Acme, Inc., Beta, Ltd.": ["Acme, Inc.", "Beta, Ltd."],
            '"Health and Human Services", "Alpha, Beta and Gamma"':
                ['"Health and Human Services"', '"Alpha, Beta and Gamma"'],
            "Alpha (East, West), Beta [North, South]": ["Alpha (East, West)", "Beta [North, South]"],
            "Alpha, Beta, and Gamma": ["Alpha", "Beta", "Gamma"],
            "- Health and Human Services\n- Smith, Jones and Co.":
                ["Health and Human Services", "Smith, Jones and Co."],
            "1. October 5, 2026\n2. October 6, 2026": ["October 5, 2026", "October 6, 2026"],
        }
        for answer, expected in cases.items():
            with self.subTest(answer=answer):
                self.assertEqual(cq.list_answer_items(answer), (expected, False))
        self.assertEqual(cq.list_answer_items("Department of Health and Human Services"),
                         (["Department of Health and Human Services"], True))
        unclear = GOOD.replace("answer_type: yes-no; number", "answer_type: list").replace(
            "Missed, by $2.8 million", "Alpha and Beta and Gamma and Delta")
        self.assertIn("list-items", self.codes(unclear, "WARN"))
        self.assertNotIn("list-contract", self.codes(unclear, "WARN"))
        # A scalar or multipart answer keeps its existing treatment, including decimal commas.
        self.assertEqual(cq.answer_items("October 5, 2026", False), ["October 5, 2026"])
        self.assertEqual(cq.answer_items("Missed, by $1,000", False), ["Missed, by $1,000"])
        self.write_platform_numbers()
        self.assertEqual(self.codes(GOOD) | self.codes(GOOD, "WARN"), set())

    def test_targeted_failure_modes_follow_the_form(self):
        # The form: 1-3 traps from its list, each once, a cause from its list (2026-10-07), then one
        # shared "Why these are traps" reason of at most 50 words (2026-10-08, replacing three boxes
        # per mode); Submit stays blocked without them or above 50 words.
        self.write_platform_numbers()
        self.assertEqual(self.codes(GOOD) | self.codes(GOOD, "WARN"), set())
        section = GOOD[GOOD.index("### Failure mode 1"):GOOD.index("## Lever checks")]
        missing = GOOD.replace(section, "")
        self.assertIn("failure-modes", self.codes(missing, "WARN"))
        self.assertNotIn("failure-modes", self.codes(missing))
        self.assertIn("failure-modes", self.codes(missing.replace("status: draft", "status: checked"), "WARN"))
        self.assertIn("failure-modes", self.codes(missing.replace("status: draft", "status: ready")))
        self.assertIn("failure-modes", self.codes(GOOD.replace("Trap: Figure", "Trap: TODO Figure"), "WARN"))
        frozen = missing.replace("status: draft", "status: submitted")
        self.assertNotIn("failure-modes", self.codes(frozen) | self.codes(frozen, "WARN"))
        modes, reason = cq.failure_section(section)
        self.assertEqual(modes, cq.failure_modes(section))
        self.assertEqual([(m["trap"], m["cause"]) for m in modes],
                         [("Figure that must be computed", "Misapplies a rule or premise")])
        self.assertTrue(reason.startswith("The second-quarter release (Doc 1) calls revenue"))
        self.assertTrue(reason.endswith("revenue missed by 2.8 million dollars."))
        self.assertEqual(cq.reason_words(reason), 42)
        # The reason may also sit under its own heading after the modes.
        heading = section.replace("Why these are traps: ", "### Why these are traps\n")
        self.assertEqual(cq.failure_section(heading), (modes, reason))
        self.assertEqual(self.codes(GOOD.replace(section, heading)) | self.codes(GOOD.replace(section, heading), "WARN"),
                         set())

        def block(number: int, trap: str) -> str:
            return (f"### Failure mode {number}\nTrap: {trap}\nHow a careless solver fails: Misapplies a rule or "
                    "premise\n\n")

        why = f"Why these are traps: {reason}\n\n"
        traps = ["Figure that must be computed", "Part vs total", "Unit or basis", "Look-alike measure"]
        failing = {
            "trap not on the list": GOOD.replace("Trap: Figure that must be computed", "Trap: Exact table vs rounded prose"),
            "cause not on the list": GOOD.replace("Misapplies a rule or premise", "Misreads the table"),
            "same trap twice": GOOD.replace(section, block(1, traps[0]) + block(2, traps[0]) + why),
            "four modes": GOOD.replace(section, "".join(block(n, trap) for n, trap in enumerate(traps, 1)) + why),
            "no shared reason": GOOD.replace(section, block(1, traps[0])),
            "reason over 50 words": GOOD.replace("2.8 million dollars.", "2.8 million dollars." + " more" * 9),
            "old per-mode box": GOOD.replace("premise\n\nWhy these", "premise\nWhere the trap sits: Doc 1, the "
                                                                  "revenue table.\n\nWhy these"),
        }
        for name, text in failing.items():
            with self.subTest(name=name):
                self.assertIn("failure-modes", self.codes(text))
        # Fifty words is the limit, the form's own option values are accepted, and a reason naming no
        # document warns.
        fifty = GOOD.replace("2.8 million dollars.", "2.8 million dollars." + " more" * 8)
        self.assertEqual(self.codes(fifty) | self.codes(fifty, "WARN"), set())
        two = GOOD.replace(section, block(1, traps[0]) + block(2, traps[2]) + why)
        self.assertEqual(self.codes(two) | self.codes(two, "WARN"), set())
        # Submit checks a pick against its definition (aq-02208 q02, 2026-10-09): the period-subtraction trap
        # wants a period in the question, and "Part vs total" a reason naming the total or the part.
        no_period = GOOD.replace("second calendar quarter of 2025", "spring 2025 season").replace(
            "for that quarter", "for that season")
        self.assertIn("failure-modes", self.codes(no_period, "WARN"))
        self.assertNotIn("failure-modes", self.codes(no_period))
        part = GOOD.replace(section, block(1, traps[1]) + why)
        self.assertIn("failure-modes", self.codes(part, "WARN"))
        part_named = part.replace("revenue missed by 2.8 million dollars.",
                                  "total revenue missed by 2.8 million dollars.")
        self.assertEqual(self.codes(part_named) | self.codes(part_named, "WARN"), set())
        values = GOOD.replace("Trap: Figure that must be computed", "Trap: derived_or_computed_metric").replace(
            "Misapplies a rule or premise", "rule_or_premise_misapplication")
        self.assertEqual(self.codes(values) | self.codes(values, "WARN"), set())
        nodoc = GOOD.replace(" (Doc 1)", "")
        self.assertIn("failure-modes", self.codes(nodoc, "WARN"))
        self.assertNotIn("failure-modes", self.codes(nodoc))

    def test_necessary_document_roles_and_alternatives(self):
        self.write_platform_numbers()
        roles = GOOD.replace("## Targeted failure modes", "## Necessary documents\n\n- D01: Answer component\n"
                                                           "- D02: Used to identify\n\n## Targeted failure modes")
        self.assertEqual(self.codes(roles) | self.codes(roles, "WARN"), set())
        self.assertIn("documents", self.codes(roles.replace("- D02: Used to identify", "- D03: Used to identify")))
        self.assertIn("documents", self.codes(roles.replace("Used to identify", "Supporting source")))
        self.assertIn("documents", self.codes(roles.replace("- D02: Used to identify", "- D02 Used to identify")))
        exclusion = roles.replace("Used to identify", "Exclusion / condition source")
        self.assertIn("documents", self.codes(exclusion, "WARN"))
        noted = exclusion.replace("condition source", "condition source | note: Rules out the revenue figure "
                                                    "restated in a later table")
        self.assertEqual(len([m for lvl, c, m in cq.check(self.write(exclusion), cq.Bundle(self.bundle))[2]
                              if c == "documents"]), 2)
        self.assertEqual(len([m for lvl, c, m in cq.check(self.write(noted), cq.Bundle(self.bundle))[2]
                              if c == "documents"]), 1)
        # Alternatives ("Doc 2 or Doc 3") are one necessary slot; each should be quoted.
        alternatives = roles.replace("documents: D01, D02", "documents: D01, D02, D03").replace(
            "- D02: Used to identify", "- D02 or D03: Answer component")
        self.assertNotIn("evidence", self.codes(alternatives))
        self.assertIn("evidence", self.codes(alternatives, "WARN"))
        quoted = alternatives.replace('"Q2 2025 | 409.8"\n', '"Q2 2025 | 409.8"\n- D03 (note) [c2]: '
                                                             '"Operating margin was 14.6% in 2025."\n')
        self.assertNotIn("evidence", self.codes(quoted) | self.codes(quoted, "WARN"))
        one_slot = roles.replace("- D01: Answer component\n- D02: Used to identify", "- D01 or D02: Answer component")
        self.assertIn("documents", self.codes(one_slot))

    def test_fingerprint_binds_roles_but_not_failure_modes(self):
        meta, sections = cq.parse(self.write(GOOD))
        legacy = {"metadata": {key: meta.get(key, "").strip() for key in
                               ("documents", "tags", "difficulty_target", "asks", "answer_type")},
                  "sections": {key: sections.get(key, "").strip() for key in
                               ("question", "answer", "components", "reasoning", "evidence")}}
        canonical = json.dumps(legacy, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
        import hashlib
        self.assertEqual(cq.submission_sha256(meta, sections), hashlib.sha256(canonical.encode("utf-8")).hexdigest())
        reworded = GOOD.replace("so a careless solver answers", "so a hasty solver answers")
        self.assertNotEqual(reworded, GOOD)
        self.assertEqual(cq.submission_sha256(*cq.parse(self.write(reworded))), cq.submission_sha256(meta, sections))
        roles = GOOD.replace("## Targeted failure modes", "## Necessary documents\n\n- D01: Answer component\n"
                                                           "- D02: Used to identify\n\n## Targeted failure modes")
        self.assertNotEqual(cq.submission_sha256(*cq.parse(self.write(roles))), cq.submission_sha256(meta, sections))

    def test_handover_writes_every_form_field_in_order(self):
        # The page numbers documents in its own order; intake.json records it for each source.
        self.write_platform_numbers()
        source = self.write(GOOD)
        with contextlib.redirect_stdout(io.StringIO()) as out:
            self.assertEqual(cq.handover(self.bundle, ["q01"]), 0)
        self.assertIn("0 FAIL, 0 WARN", out.getvalue())
        json_path = source.with_suffix(".json")
        self.assertEqual(json.loads(json_path.read_bytes())["gold_answer"], "Missed, by $2.8 million")
        self.assertEqual(source.read_bytes(), GOOD.encode("utf-8"))
        sheet = (self.bundle / "notes" / "handover-q01.md").read_bytes().decode("utf-8")
        headings = ["### Question", "### Number of asks", "### Gold answer", "### Step-by-step reasoning",
                    "### Answer components", "### Evidence", "### Necessary documents", "### Reasoning tags",
                    "### Targeted failure modes"]
        self.assertEqual([sheet.index(h) for h in headings], sorted(sheet.index(h) for h in headings))
        self.assertIn("D01 = Doc 2, D02 = Doc 1", sheet)
        # Wrapped lines become one paragraph in a box; the reasoning keeps its lines as steps.
        self.assertIn("```text\nDid the freight company's actual revenue for the second calendar quarter of 2025 miss "
                      "or beat the revenue it had forecast", sheet)
        self.assertIn('#### Why these are traps (at most 50 words)\n\n```text\nThe second-quarter release (Doc 1) calls '
                      'revenue "in line" with the forecast, so a careless solver answers', sheet)
        self.assertIn("is in\nthe second-quarter table.", sheet)
        self.assertIn("### Number of asks\n\n2 asks", sheet)
        self.assertIn("Evidence 4: document Doc 1 (b-release-q2.html), proves component 2. $2.8 million", sheet)
        self.assertIn("- Tick Doc 2 (a-release-q1.pdf), role: Answer component", sheet)
        self.assertIn("Leave unticked: Doc 3 (c-notes.txt)", sheet)
        self.assertIn("Tick: Cross-document join, Temporal alignment\n\nLeave unticked: Unit normalization", sheet)
        self.assertIn("How a careless solver fails: Misapplies a rule or premise", sheet)
        self.assertIn("(42 / 50 words.)", sheet)
        self.assertNotIn("Where the trap sits", sheet)
        self.assertIn("DRAFT", sheet)
        self.assertIn("so are the default Answer component roles shown below", sheet)
        self.assertNotIn("Authorship", sheet)
        self.assertNotIn("Lever checks", sheet)
        claude = GOOD.replace("origin: user", "origin: claude").replace("status: draft", "status: checked")
        self.write(claude)
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cq.handover(self.bundle, ["q01"]), 0)
        sheet = (self.bundle / "notes" / "handover-q01.md").read_bytes().decode("utf-8")
        self.assertIn("Authorship: origin claude", sheet)
        self.assertNotIn("DRAFT", sheet)
        self.assertEqual(source.read_bytes(), claude.encode("utf-8"))
        self.assertNotIn("origin", json.loads(json_path.read_bytes()))
        json_path.unlink()
        self.write(GOOD.replace("Missed, by $2.8 million", "TODO"))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cq.handover(self.bundle, ["q01"]), 1)
        self.assertIn("NOT READY TO PASTE", (self.bundle / "notes" / "handover-q01.md").read_bytes().decode("utf-8"))
        self.assertFalse(json_path.exists())
        (self.bundle / "notes" / "handover-q01.md").unlink()
        self.write(GOOD.replace("status: draft", "status: submitted"))
        with contextlib.redirect_stdout(io.StringIO()) as out:
            cq.handover(self.bundle, ["q01"])
        self.assertIn("no hand-over", out.getvalue())
        self.assertFalse((self.bundle / "notes" / "handover-q01.md").exists())
        self.assertFalse(json_path.exists())

    def test_json_inputs_contain_only_platform_fields_and_expand_evidence(self):
        self.write_platform_numbers()
        text = GOOD.replace("origin: user", "origin: claude").replace(
            "status: draft", "status: checked").replace("The forecast of", "Le prévisionnel — the forecast of")
        text += "\n## Platform validation\n\nPRIVATE PLATFORM RETURN\n\n## Local resolution\n\nPRIVATE DIAGNOSTIC\n"
        source = self.write(text)
        meta, sec = cq.parse(source)
        bundle = cq.Bundle(self.bundle)
        with mock.patch.object(cq, "check", side_effect=AssertionError("serialization must not run checks")):
            target = cq.write_json_inputs(bundle, source, meta, sec)
        self.assertEqual(target, source.with_suffix(".json"))
        self.assertEqual(source.read_bytes(), text.encode("utf-8"))
        data = json.loads(target.read_bytes())
        self.assertEqual(set(data), {"question", "number_of_asks", "gold_answer", "step_by_step_reasoning",
                                     "answer_components", "evidence", "necessary_documents", "reasoning_tags",
                                     "targeted_failure_modes", "why_these_are_traps"})
        self.assertEqual(data["question"], "Did the freight company's actual revenue for the second calendar "
                         "quarter of 2025 miss or beat the revenue it had forecast for that quarter, and by how "
                         "many million dollars, to one decimal place?")
        self.assertEqual(data["number_of_asks"], 2)
        self.assertIs(type(data["number_of_asks"]), int)
        self.assertEqual(data["gold_answer"], "Missed, by $2.8 million")
        self.assertEqual(data["step_by_step_reasoning"], sec["reasoning"])
        self.assertIn("prévisionnel —", data["step_by_step_reasoning"])
        self.assertEqual(data["answer_components"], ["Missed", "$2.8 million"])
        self.assertEqual(data["evidence"], [
            {"document_number": 2, "component_number": 1, "quote": "expects second-quarter revenue of $412.6 million"},
            {"document_number": 2, "component_number": 2, "quote": "expects second-quarter revenue of $412.6 million"},
            {"document_number": 1, "component_number": 1, "quote": "Q2 2025 | 409.8"},
            {"document_number": 1, "component_number": 2, "quote": "Q2 2025 | 409.8"}])
        self.assertEqual(data["necessary_documents"], [
            {"document_numbers": [2], "role": "Answer component", "reviewer_note": ""},
            {"document_numbers": [1], "role": "Answer component", "reviewer_note": ""}])
        self.assertEqual(data["reasoning_tags"], ["Cross-document join", "Temporal alignment"])
        self.assertEqual(len(data["targeted_failure_modes"]), 1)
        mode = data["targeted_failure_modes"][0]
        self.assertEqual(mode, {"trap": "Figure that must be computed",
                                "how_a_careless_solver_fails": "Misapplies a rule or premise"})
        self.assertTrue(data["why_these_are_traps"].startswith('The second-quarter release (Doc 1) calls revenue "in line"'))
        self.assertTrue(data["why_these_are_traps"].endswith("revenue missed by 2.8 million dollars."))
        self.assertNotIn("\n", data["why_these_are_traps"])
        self.assertNotIn("PRIVATE", target.read_bytes().decode("utf-8"))
        self.assertNotIn("One document is not enough", target.read_bytes().decode("utf-8"))

    def test_json_inputs_preserve_component_and_alternative_order_and_normalize_labels(self):
        self.write_platform_numbers()
        text = GOOD.replace("c1", "c7").replace("c2", "c3").replace(
            "documents: D01, D02", "documents: D01, D02, D03").replace(
            "tags: Cross-document join; Temporal alignment", "tags: temporal alignment; cross-document join").replace(
            "Trap: Figure that must be computed", "Trap: derived_or_computed_metric").replace(
            "How a careless solver fails: Misapplies a rule or premise",
            "How a careless solver fails: rule_or_premise_misapplication").replace(
            "## Targeted failure modes", "## Necessary documents\n\n- D01: answer component\n"
            "- D03 or D02: exclusion / condition source | note: Both releases use the same fiscal period.\n\n"
            "## Targeted failure modes")
        meta, sec = cq.parse(self.write(text))
        data = cq.form_payload(meta, sec, cq.Bundle(self.bundle))
        self.assertEqual(data["answer_components"], ["Missed", "$2.8 million"])
        self.assertEqual([row["component_number"] for row in data["evidence"]], [1, 2, 1, 2])
        self.assertEqual(data["necessary_documents"], [
            {"document_numbers": [2], "role": "Answer component", "reviewer_note": ""},
            {"document_numbers": [3, 1], "role": "Exclusion / condition source",
             "reviewer_note": "Both releases use the same fiscal period."}])
        self.assertEqual(set(data["reasoning_tags"]), {"Cross-document join", "Temporal alignment"})
        self.assertEqual(data["targeted_failure_modes"][0]["trap"], "Figure that must be computed")
        self.assertEqual(data["targeted_failure_modes"][0]["how_a_careless_solver_fails"], "Misapplies a rule or premise")

    def test_json_inputs_refuse_unknown_mapping_and_malformed_required_fields(self):
        source = self.write(GOOD)
        meta, sec = cq.parse(source)
        with self.assertRaises(ValueError):
            cq.write_json_inputs(cq.Bundle(self.bundle), source, meta, sec)
        self.assertFalse(source.with_suffix(".json").exists())
        self.assertEqual(source.read_bytes(), GOOD.encode("utf-8"))
        self.write_platform_numbers()
        bad_inputs = {
            "unknown component": GOOD.replace("[c1, c2]", "[c1, c9]"),
            "unknown evidence document": GOOD.replace("- D02 (revenue table)", "- D04 (revenue table)"),
            "unparsed evidence": GOOD.replace('- D02 (revenue table) [c1, c2]: "Q2 2025 | 409.8"',
                                               "- this evidence line cannot be parsed"),
            "unparsed components": GOOD.replace("- c2: $2.8 million", "- this component cannot be parsed"),
            "unknown tag": GOOD.replace("tags: Cross-document join; Temporal alignment", "tags: Unknown tag"),
            "malformed document metadata": GOOD.replace("documents: D01, D02", "documents: D01, unknown"),
            "unknown role": GOOD.replace("## Targeted failure modes", "## Necessary documents\n\n"
                                          "- D01: unknown role\n- D02: Answer component\n\n## Targeted failure modes"),
            "unknown trap": GOOD.replace("Trap: Figure that must be computed", "Trap: unknown trap"),
            "unknown cause": GOOD.replace("How a careless solver fails: Misapplies a rule or premise",
                                            "How a careless solver fails: unknown cause"),
            "missing asks": GOOD.replace("asks: 2\n", ""),
            "missing question": GOOD.replace("## Question", "## Local question"),
            "missing answer": GOOD.replace("## Answer", "## Local answer"),
            "missing reasoning": GOOD.replace("## Reasoning", "## Local reasoning"),
            "missing failure modes": GOOD.replace("## Targeted failure modes", "## Local failure modes"),
            "missing failure-mode reason": GOOD.replace("Why these are traps: ", "Notes: "),
            "failure-mode reason over 50 words": GOOD.replace("2.8 million dollars.", "2.8 million dollars." + " more" * 9),
            "old per-mode failure box": GOOD.replace("premise\n\nWhy these", "premise\nWhere the trap sits: Doc 1.\n\n"
                                                                          "Why these"),
        }
        for name, text in bad_inputs.items():
            with self.subTest(name=name):
                source = self.write(text)
                meta, sec = cq.parse(source)
                with self.assertRaises(ValueError):
                    cq.write_json_inputs(cq.Bundle(self.bundle), source, meta, sec)
                self.assertFalse(source.with_suffix(".json").exists())
                self.assertEqual(source.read_bytes(), text.encode("utf-8"))

    def test_json_cli_serializes_selected_frozen_question_without_validation(self):
        self.write_platform_numbers()
        text = GOOD.replace("origin: user", "origin: claude").replace("status: draft", "status: accepted")
        source = self.write(text)
        self.write(GOOD.replace("id: q01", "id: q02"), "q02")
        with mock.patch.object(cq, "check", side_effect=AssertionError("JSON must not run the floor")), \
                mock.patch.object(cq.pc, "summary", side_effect=AssertionError("JSON must not consult platform checks")), \
                contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cq.main([str(self.bundle), "--json", "--q", "q01"]), 0)
        self.assertEqual(source.read_bytes(), text.encode("utf-8"))
        self.assertTrue(source.with_suffix(".json").exists())
        self.assertFalse((self.bundle / "questions" / "q02.json").exists())
        self.assertFalse((self.bundle / "notes" / "handover-q01.md").exists())

    def test_handover_reports_missing_json_mapping_without_changing_source(self):
        source = self.write(GOOD)
        with contextlib.redirect_stdout(io.StringIO()) as out:
            self.assertEqual(cq.handover(self.bundle, ["q01"]), 0)
        self.assertIn("JSON not written", out.getvalue())
        self.assertTrue((self.bundle / "notes" / "handover-q01.md").exists())
        self.assertFalse(source.with_suffix(".json").exists())
        self.assertEqual(source.read_bytes(), GOOD.encode("utf-8"))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cq.main([str(self.bundle), "--json", "--q", "q01"]), 1)
        self.assertFalse(source.with_suffix(".json").exists())

    def test_sibling_questions_on_one_bundle(self):
        self.write(GOOD.replace("id: q01", "id: q02"), "q02")
        self.assertIn("sibling", self.codes(GOOD, "WARN"))
        self.write(GOOD.replace("id: q01", "id: q02").replace("status: draft", "status: dropped"), "q02")
        self.assertNotIn("sibling", self.codes(GOOD, "WARN"))

    def test_legacy_target_needs_bound_hard_platform_pass_and_export_takes_only_ready(self):
        ready = self.rebound_ready(self.targeted_ready("hard", local=False).replace("difficulty_target: hard\n", ""))
        self.assertEqual(self.codes(ready), set())
        failing = {
            "platform solved it": ready.replace("platform_verdict: pass", "result: solved | platform_verdict: pass"),
            "no platform result": re.sub(r"(?m)^- .*platform difficulty check.*\n", "", ready),
            "High tier not reached": ready.replace("High 0/5", "High —"),
            "below legacy target": ready.replace("passed_tier: hard", "passed_tier: easy"),
            "observed counts are not a pass": ready.replace("platform_verdict: pass | ", ""),
            "no passed tier": ready.replace("passed_tier: hard | ", ""),
            "no saved platform feedback": ready.replace("feedback: Passed hard verification", ""),
            "contradictory platform verdict": ready.replace("Passed hard verification", "Too easy: High model solved it"),
            "records of an older revision": ready.replace("revision: 0", "revision: 1"),
            "untagged records": ready.replace("rev: 0 | ", ""),
            "unbound historical records": re.sub(r"tested_sha256: [a-f0-9]+ \| ", "", ready),
        }
        for name, text in failing.items():
            with self.subTest(name=name):
                self.assertIn("status", self.codes(text))
        self.write(GOOD.replace("id: q01", "id: q02"), "q02")
        self.write(ready.replace("id: q01", "id: q03").replace("origin: user", "origin: claude"), "q03")
        self.write(ready)
        with contextlib.redirect_stdout(io.StringIO()) as out:
            self.assertEqual(cq.export(self.bundle, self.root), 0)
        sheet = (self.root / "dist" / "freight.md").read_bytes().decode("utf-8")
        self.assertIn("## q01", sheet)
        self.assertNotIn("## q02", sheet)
        self.assertNotIn("## q03", sheet)
        self.assertIn("skipped q03", out.getvalue())
        self.assertIn("### Gold answer\n\n```text\nMissed, by $2.8 million\n```", sheet)
        self.assertIn("### Answer components\n\nComponent 1\n\n```text\nMissed\n```\n\nComponent 2\n\n"
                      "```text\n$2.8 million\n```", sheet)
        self.assertIn("proves component 2. $2.8 million (local note: revenue table)\n\n```text\nQ2 2025 | 409.8\n```",
                      sheet)
        self.assertIn("- Tick D02 (b-release-q2.html; page number unknown), role: Answer component", sheet)
        self.assertIn("Trap: Figure that must be computed", sheet)
        self.assertNotIn("Lever checks", sheet)
        self.assertNotIn("outcome:", sheet)

    def targeted_ready(self, target: str = "easy", passed_tier: str | None = None, local: bool = True) -> str:
        """Build a complete revision-bound fixture; this is a recorded verdict, not a live pass."""
        # The extraction fixture intentionally includes one unsupported document; a complete
        # corpus for readiness excludes that separate unsupported-format test input.
        manifest_path = self.bundle / "text" / "manifest.json"
        manifest = json.loads(manifest_path.read_bytes())
        if "D05" in manifest:
            manifest.pop("D05")
            manifest_path.write_bytes(json.dumps(manifest).encode())
            (self.bundle / "docs" / "e-sheet.xlsx").unlink()
        passed_tier = passed_tier or target
        question = GOOD.replace("status: draft", "status: ready").replace(
            "revision: 0", f"revision: 0\ndifficulty_target: {target}")
        meta, sections = cq.parse(self.write(question))
        digest = cq.submission_sha256(meta, sections, cq.Bundle(self.bundle))
        binding = f"rev: 0 | tested_sha256: {digest}"
        tier_line = {"easy": "Easy 1/5 · Medium — · Hard —",
                     "medium": "Easy 1/5 · Medium 0/5 · Hard —",
                     "hard": "Low 1/5 · Medium 0/5 · High 0/5"}[passed_tier]
        platform = (f"- 2026-10-02 | {binding} | platform difficulty check | platform_verdict: pass | "
                    f"passed_tier: {passed_tier} | tiers: {tier_line} | feedback: Passed {passed_tier} verification\n")
        if not local:
            return question + "\n## Platform validation\n\n" + platform
        solves = (f"- 2026-10-01 | {binding} | solver: A | outcome: match | answer: below by 2.8 | readings: none\n"
                  f"- 2026-10-01 | {binding} | solver: B | outcome: match | answer: below by 2.8 | readings: none\n"
                  f"- 2026-10-01 | {binding} | closed-book solve | result: misses | answer: unknown\n"
                  + platform)
        review = f"- 2026-10-01 | {binding} | verdict: ready | difficulty: {target} | findings: none\n"
        audit = f"## Shortcut audit\n\n- 2026-10-01 | {binding} | verdict: clear | findings: none\n\n"
        return question.replace("## Blind solve\n", "## Blind solve\n\n" + solves).replace(
            "## Review\n", "## Review\n\n" + review).replace("## Feedback log\n", audit + "## Feedback log\n")

    def test_targeted_ready_uses_declared_platform_pass_not_observed_tier(self):
        for target in ("easy", "medium", "hard"):
            with self.subTest(target=target):
                self.assertEqual(self.codes(self.targeted_ready(target)), set())
        self.assertEqual(self.codes(self.targeted_ready("easy", "hard")), set())
        self.assertIn("status", self.codes(self.targeted_ready("hard", "easy")))
        ready = self.targeted_ready()
        cases = {
            "unknown verdict": ready.replace("platform_verdict: pass", "platform_verdict: unknown"),
            "explicit failure": ready.replace("platform_verdict: pass", "platform_verdict: fail"),
            "numeric output alone": ready.replace("platform_verdict: pass | passed_tier: easy | ", ""),
            "tier did not run": ready.replace("Easy 1/5", "Easy —"),
            "invalid attempt count": ready.replace("Easy 1/5", "Easy 6/5"),
            "zero attempts": ready.replace("Easy 1/5", "Easy 0/0"),
            "missing feedback": ready.replace("feedback: Passed easy verification", "feedback:"),
            "contradictory feedback": ready.replace("Passed easy verification", "Too easy — solver got it"),
            "unfinished platform feedback": ready.replace("Passed easy verification", "Pending / Running / Not started"),
            "rejection hidden after delimiter": ready.replace("Passed easy verification", "Passed easy verification | Too easy"),
            "conflicting verdict fields": ready.replace("platform_verdict: pass", "platform_verdict: fail | platform_verdict: pass"),
            "shortcut found": ready.replace("verdict: clear", "verdict: shortcut"),
        }
        for name, question in cases.items():
            with self.subTest(name=name):
                self.assertIn("status", self.codes(question))
        # Counts never manufacture a verdict or an undocumented threshold.
        self.assertEqual(self.codes(ready.replace("Easy 1/5", "Easy 5/5")), set())
        self.assertEqual(self.codes(ready.replace("Passed easy verification", "Easy passed; Medium failed")), set())
        self.assertEqual(self.codes(ready.replace("Passed easy verification", "Easy passed; Medium too easy")), set())
        self.assertEqual(self.codes(ready.replace("Passed easy verification", "Easy passed; Medium pending")), set())
        hard = self.targeted_ready("hard")
        self.assertIn("status", self.codes(hard.replace("Passed hard verification", "Easy too easy; Hard passed")))
        self.assertIn("status", self.codes(hard.replace("Passed hard verification", "Low failed; High passed")))
        for unfinished in ("Pending", "Running", "Not started", "In progress", "Queued", "Easy verification pending",
                           "No verdict yet", "Not yet tested", "Awaiting results"):
            with self.subTest(unfinished=unfinished):
                self.assertIn("status", self.codes(ready.replace("Passed easy verification", unfinished)))
        self.assertIn("status", self.codes(ready.replace("Passed easy verification", "Easy failed; Medium passed")))
        self.assertIn("difficulty-target", self.codes(ready.replace("difficulty_target: easy", "difficulty_target: low")))

    def test_local_positives_predictions_and_freshness_are_not_ready_requirements(self):
        ready = self.targeted_ready()
        cases = {
            "platform only": self.targeted_ready(local=False),
            "one gold match": ready.replace("outcome: match", "outcome: wrong-lever", 1),
            "no gold matches": ready.replace("outcome: match", "outcome: solver-error"),
            "repeated solver": ready.replace("solver: B", "solver: A"),
            "missing solver identity": ready.replace("solver: B | ", ""),
            "no closed-book": re.sub(r"(?m)^- .*closed-book solve.*\n", "", ready),
            "no audit": ready.replace("## Shortcut audit", "## Historical optional audit"),
            "no review": ready.replace("## Review", "## Historical optional review"),
            "lower local rating": ready.replace("difficulty: easy", "difficulty: below maximum"),
            "no local rating": ready.replace("difficulty: easy | ", ""),
            "stale local records": re.sub(r"(?m)^(- 2026-10-01 .*tested_sha256: )[a-f0-9]+", r"\g<1>" + "0" * 64, ready),
            "unbound newest review": ready.replace("## Shortcut audit", "- 2026-10-03 | rev: 0 | verdict: ready\n\n## Shortcut audit"),
        }
        for name, question in cases.items():
            with self.subTest(name=name):
                self.assertEqual(self.codes(question), set())
        # A generic local 'revise' is visible, but never substitutes its rating for the platform.
        diagnostic = ready.replace("verdict: ready", "verdict: revise").replace(
            "findings: none", "findings: The local rating may be below Easy.", 1)
        self.assertEqual(self.codes(diagnostic), set())
        _, _, found = cq.check(self.write(diagnostic), cq.Bundle(self.bundle))
        self.assertTrue(any(level == "WARN" and code == "local-diagnostic" and "below Easy" in message
                            for level, code, message in found))

    def test_current_adverse_local_evidence_survives_later_positive_records(self):
        ready = self.targeted_ready()
        binding = re.search(r"rev: 0 \| tested_sha256: [a-f0-9]+", ready).group()
        cases = {
            "ambiguity": ready.replace("outcome: match", "outcome: ambiguous", 1),
            "unsupported gold": ready.replace("outcome: match", "outcome: not-answerable", 1),
            "closed-book reach": ready.replace("result: misses", "result: reaches"),
            "removal reach": ready.replace("## Blind solve\n", "## Blind solve\n\n"
                f"- 2026-10-01 | {binding} | removal probe | left_out: D01 | result: reaches\n"),
            "shortcut": ready.replace("verdict: clear", "verdict: shortcut"),
            "review unsupported": ready.replace("verdict: ready", "verdict: unsupported"),
            "invalid outcome": ready.replace("outcome: match", "outcome: wronglever", 1),
            "duplicate outcome": ready.replace("outcome: match", "outcome: ambiguous | outcome: match", 1),
        }
        for name, question in cases.items():
            with self.subTest(name=name):
                self.assertIn("status", self.codes(question))
                # Historical diagnostics remain in the file without forcing local reruns.
                historical = question.replace("- 2026-10-01 | rev: 0", "- 2026-10-01 | rev: 1")
                self.assertEqual(self.codes(historical), set())
                unbound = re.sub(r"(?m)^(- 2026-10-01 .*tested_sha256: )[a-f0-9]+", r"\g<1>" + "0" * 64, question)
                self.assertEqual(self.codes(unbound), set())

    def test_concrete_local_resolution_needs_attribution_and_current_binding(self):
        ready = self.targeted_ready().replace("outcome: match", "outcome: ambiguous", 1)
        binding = re.search(r"rev: 0 \| tested_sha256: [a-f0-9]+", ready).group()
        resolution = (f"- 2026-10-03 | {binding} | assessor: Author | issue: blind: ambiguous | "
                      "verdict: resolved | findings: D01 labels the forecast for Q2; the alternative used "
                      "the Q1 actual, which the exact question period excludes.\n")
        resolved = ready + "\n## Local resolution\n\n" + resolution
        self.assertEqual(self.codes(resolved), set())
        _, _, found = cq.check(self.write(resolved), cq.Bundle(self.bundle))
        self.assertTrue(any(code == "local-diagnostic" and "adverse history retained" in message
                            for _, code, message in found))
        for name, bad in {
            "no assessor": resolution.replace("assessor: Author | ", ""),
            "wrong issue": resolution.replace("issue: blind: ambiguous", "issue: closed-book: reaches"),
            "older revision": resolution.replace("rev: 0", "rev: 1"),
            "stale fingerprint": re.sub(r"tested_sha256: [a-f0-9]+", "tested_sha256: " + "0" * 64, resolution),
            "not resolved": resolution.replace("verdict: resolved", "verdict: pending"),
            "duplicate verdict": resolution.replace("verdict: resolved", "verdict: pending | verdict: resolved"),
            "missing findings": resolution.split("findings:")[0],
        }.items():
            with self.subTest(name=name):
                self.assertIn("status", self.codes(ready + "\n## Local resolution\n\n" + bad))
        self.assertIn("status", self.codes(resolved + resolution.replace("verdict: resolved", "verdict: pending")))

    def test_platform_history_merges_sections_and_never_hides_a_later_failure(self):
        ready = self.targeted_ready(local=False)
        platform = next(line for line in ready.splitlines() if "platform difficulty check" in line)
        self.assertEqual(self.codes(ready), set())
        self.assertEqual(self.codes(ready.replace("platform difficulty check", "platform validation")), set())
        legacy_failure = platform.replace("2026-10-02", "2026-10-03").replace("platform_verdict: pass", "platform_verdict: fail")
        both = ready.replace("## Blind solve\n", "## Blind solve\n\n" + legacy_failure + "\n")
        _, sec = cq.parse(self.write(both))
        self.assertEqual([r["platform_verdict"] for r in cq.platform_records(sec)], ["pass", "fail"])
        self.assertIn("status", self.codes(both))
        # Dedicated-section precedence only breaks a date tie; it cannot hide newer legacy feedback.
        self.assertEqual(self.codes(both.replace("2026-10-03", "2026-10-01")), set())
        self.assertEqual(self.codes(both.replace("2026-10-03", "2026-10-02")), set())
        later_unbound = re.sub(r"tested_sha256: [a-f0-9]+", "tested_sha256: " + "0" * 64,
                               platform.replace("2026-10-02", "2026-10-04"))
        self.assertIn("status", self.codes(ready + later_unbound + "\n"))
        for unknown_date in ("today", "2026-99-99"):
            with self.subTest(unknown_date=unknown_date):
                undated = both.replace("2026-10-03", unknown_date)
                self.assertIn("status", self.codes(undated))
                self.assertIn("status", self.codes(ready.replace("2026-10-02", unknown_date)))

    def test_targeted_check_records_bind_to_submission_and_corpus(self):
        ready = self.targeted_ready()
        self.assertEqual(self.codes(ready), set())
        for name, question in {
            "question edited without revision bump": ready.replace("actual revenue", "reported actual revenue", 1),
            "reasoning edited": ready.replace("forecast of $412.6", "original forecast of $412.6"),
            "evidence location edited": ready.replace("D01 (p. 1)", "D01 (first release)"),
            "hash missing from records": re.sub(r"tested_sha256: [a-f0-9]+ \| ", "", ready),
            "older revision": ready.replace("revision: 0", "revision: 1"),
        }.items():
            with self.subTest(name=name):
                self.assertIn("status", self.codes(question))
        for path in (self.bundle / "text" / "D01.txt", self.bundle / "docs" / "c-notes.txt",
                     self.bundle / "text" / "manifest.json"):
            original = path.read_bytes()
            try:
                path.write_bytes(original + b"\n ")
                with self.subTest(changed_corpus=path.name):
                    self.assertIn("status", self.codes(ready))
            finally:
                path.write_bytes(original)
        # Recording findings/status and advancing a workflow revision do not mutate the submission.
        meta, sections = cq.parse(self.write(ready))
        digest = cq.submission_sha256(meta, sections, cq.Bundle(self.bundle))
        meta.update(status="checked", revision="4")
        sections["review"] = "changed review findings"
        self.assertEqual(digest, cq.submission_sha256(meta, sections, cq.Bundle(self.bundle)))
        for key in ("documents", "tags", "difficulty_target", "asks", "answer_type"):
            changed = {**meta, key: meta[key] + " changed"}
            self.assertNotEqual(digest, cq.submission_sha256(changed, sections, cq.Bundle(self.bundle)))

    def test_known_incomplete_capture_cannot_be_ready(self):
        notes = self.bundle / "notes"
        notes.mkdir()
        intake = notes / "intake.json"
        for fields in ({"incomplete": True}, {"gap_characters": 8}, {"matches_capture": False}):
            intake.write_bytes(json.dumps({"schema": 1, "documents": [{"source": "docs/a-release-q1.pdf", **fields}]}).encode())
            with self.subTest(fields=fields):
                self.assertIn("source-completeness", self.codes(self.targeted_ready()))
        intake.unlink()
        text = self.bundle / "text" / "D03.txt"
        text.write_bytes(text.read_bytes() + b"\n[... 8 characters not shown on the page ...]\n")
        self.assertIn("source-completeness", self.codes(self.targeted_ready()))

    def test_missing_source_or_extraction_cannot_be_ready(self):
        ready = self.targeted_ready()
        for source in (self.bundle / "text" / "D03.txt", self.bundle / "docs" / "c-notes.txt"):
            original = source.read_bytes()
            try:
                source.unlink()
                meta, sections = cq.parse(self.write(ready))
                digest = cq.submission_sha256(meta, sections, cq.Bundle(self.bundle))
                rebound = re.sub(r"tested_sha256: [a-f0-9]+", "tested_sha256: " + digest, ready)
                with self.subTest(source=source.name):
                    self.assertIn("source-completeness", self.codes(rebound))
            finally:
                source.write_bytes(original)

    def rebound_ready(self, ready: str) -> str:
        """Bind synthetic fixture records to changed bytes; never a real check or verdict."""
        meta, sections = cq.parse(self.write(ready))
        digest = cq.submission_sha256(meta, sections, cq.Bundle(self.bundle))
        return re.sub(r"tested_sha256: [a-f0-9]+", "tested_sha256: " + digest, ready)

    def assessed_ready(self, ready: str) -> str:
        """A synthetic reviewer record tests the schema, not whether prose is true."""
        meta, sections = cq.parse(self.write(ready))
        corpus = cq.Bundle(self.bundle)
        digest = cq.submission_sha256(meta, sections, corpus)
        gaps = cq.source_completeness(corpus)
        explanation = {
            "premises": "The intact D03 original discusses an operating margin, not the Q2 revenue forecast.",
            "gold": "D01 and D02 quote the forecast and actual revenue; the D03 original supplies no revenue.",
            "exclusions": "This question compares two fixed Q2 values, rather than selecting an exhaustive list.",
            "uniqueness": "The intact D03 original contains no alternate quarter, forecast or revenue estimate.",
            "necessity": "Removing D01 loses the forecast and removing D02 loses actual revenue; D03 has neither.",
            "shortcut": "Reading the intact D03 original offers only operating margin and no comparison shortcut.",
        }
        findings = json.dumps({gap: explanation for gap in gaps}, ensure_ascii=False)
        record = (f"- 2026-10-05 | rev: 0 | reviewer: Reviewer-C | tested_sha256: {digest} | "
                  f"source_gaps_sha256: {cq.source_gaps_sha256(gaps)} | verdict: unaffected | findings: {findings}\n")
        return ready + "\n## Source assessment\n\n" + record

    def test_scoped_source_assessment_preserves_incomplete_warning_and_platform_gate(self):
        ready = self.targeted_ready(local=False)
        # The author can inspect the intact original of the unused margin-only note.
        (self.bundle / "text" / "D03.txt").unlink()
        ready = self.rebound_ready(ready)
        self.assertIn("source-completeness", self.codes(ready))
        assessed = self.assessed_ready(ready)
        self.assertEqual(self.codes(assessed), set())
        _, _, found = cq.check(self.write(assessed), cq.Bundle(self.bundle))
        self.assertTrue(any(level == "WARN" and code == "source-completeness" and
                            "corpus remains incomplete" in message for level, code, message in found))
        self.assertTrue(cq.source_completeness(cq.Bundle(self.bundle)))
        self.assertIn("status", self.codes(assessed.replace("platform_verdict: pass", "platform_verdict: unknown")))
        for identity in ("assessor: Author", "author: Codex", "assessor: Platform source review"):
            with self.subTest(identity=identity):
                self.assertEqual(self.codes(assessed.replace("reviewer: Reviewer-C", identity)), set())

    def test_source_assessment_must_cover_exact_gaps_and_all_impacts(self):
        ready = self.targeted_ready()
        (self.bundle / "text" / "D03.txt").unlink()
        assessed = self.assessed_ready(self.rebound_ready(ready))
        gap = cq.source_completeness(cq.Bundle(self.bundle))[0]
        meta, sections = cq.parse(self.write(assessed))
        fields, _, _ = cq.record_fields(sections["source assessment"].splitlines()[0])
        original = fields["findings"]
        findings = json.loads(original)
        wrong_key = {"D03: another gap": findings[gap]}
        incomplete = {gap: {key: value for key, value in findings[gap].items() if key != "shortcut"}}
        placeholder = {gap: {key: "not applicable" for key in findings[gap]}}
        for name, text in {
            "different gap identity": assessed.replace(original, json.dumps(wrong_key)),
            "missing gap": assessed.replace(original, "{}"),
            "missing impact": assessed.replace(original, json.dumps(incomplete)),
            "generic assurance": assessed.replace(original, json.dumps(placeholder)),
            "findings not JSON": assessed.replace(original, "All gaps are irrelevant to this question."),
            "duplicate finding keys": assessed.replace(original, "{" + json.dumps(gap) + ": " +
                                                      json.dumps(findings[gap]) + ", " + json.dumps(gap) +
                                                      ": " + json.dumps(findings[gap]) + "}"),
            "wrong gap hash": re.sub(r"source_gaps_sha256: [a-f0-9]+", "source_gaps_sha256: " + "0" * 64, assessed),
            "missing reviewer": assessed.replace("reviewer: Reviewer-C | ", ""),
            "affected verdict": assessed.replace("verdict: unaffected", "verdict: affected"),
            "duplicate structural fields": assessed.replace("verdict: unaffected", "verdict: affected | verdict: unaffected"),
            "wrong revision": assessed.replace("- 2026-10-05 | rev: 0", "- 2026-10-05 | rev: 1"),
            "latest assessment unbound": assessed + "- 2026-10-05 | rev: 0 | reviewer: Reviewer-C | verdict: unaffected\n",
        }.items():
            with self.subTest(name=name):
                self.assertIn("source-completeness", self.codes(text))
        # A delimiter inside a terminal JSON explanation remains prose, not typed fields.
        quoted = assessed.replace("operating margin, not", "operating margin | verdict: affected, not")
        self.assertEqual(self.codes(quoted), set())

    def test_source_assessment_rechecks_corpus_and_new_gap_identities(self):
        ready = self.targeted_ready()
        (self.bundle / "text" / "D03.txt").unlink()
        assessed = self.assessed_ready(self.rebound_ready(ready))
        self.assertEqual(self.codes(assessed), set())
        (self.bundle / "docs" / "c-notes.txt").write_bytes(b"Different source bytes.")
        self.assertIn("source-completeness", self.codes(assessed))
        # Rerunning other records cannot rescue an assessment that omits a newly known gap.
        (self.bundle / "text" / "D04.txt").unlink()
        rebound = self.rebound_ready(assessed)
        self.assertIn("source-completeness", self.codes(rebound))
        gaps = cq.source_completeness(cq.Bundle(self.bundle))
        self.assertEqual(cq.source_gaps_sha256(gaps), cq.source_gaps_sha256(list(reversed(gaps))))

    def test_source_assessment_cannot_waive_necessary_cited_or_unscoped_gaps(self):
        ready = self.targeted_ready()
        for relative in ("text/D01.txt", "docs/a-release-q1.pdf"):
            path = self.bundle / relative
            original = path.read_bytes()
            try:
                path.unlink()
                assessed = self.assessed_ready(self.rebound_ready(ready))
                with self.subTest(missing=relative):
                    self.assertIn("source-completeness", self.codes(assessed))
            finally:
                path.write_bytes(original)
        manifest_path = self.bundle / "text" / "manifest.json"
        original_manifest = manifest_path.read_bytes()
        manifest = json.loads(original_manifest)
        manifest["D01"]["method"] = "unsupported"
        manifest_path.write_bytes(json.dumps(manifest).encode())
        self.assertIn("source-completeness", self.codes(self.assessed_ready(self.rebound_ready(ready))))
        manifest_path.write_bytes(original_manifest)
        # A quote cannot escape source protection simply by omitting the necessary-doc tag.
        (self.bundle / "docs" / "c-notes.txt").unlink()
        cited = ready.replace("## Evidence\n", "## Evidence\n\n- D03 (note) [c1]: \"Operating margin was 14.6% in 2025.\"\n")
        self.assertIn("source-completeness", self.codes(self.assessed_ready(self.rebound_ready(cited))))
        (self.bundle / "notes").mkdir()
        (self.bundle / "notes" / "intake.json").write_bytes(b'{"documents": "invalid"}')
        self.assertIn("source-completeness", self.codes(self.assessed_ready(self.rebound_ready(ready))))

    def test_source_gap_identities_include_capture_size_and_marker_location(self):
        ready = self.targeted_ready()
        (self.bundle / "notes").mkdir()
        (self.bundle / "notes" / "intake.json").write_bytes(json.dumps({"documents": [{
            "source": "docs/c-notes.txt", "incomplete": True, "gap_characters": 8,
            "note": "only 32 of 40 characters shown", "matches_capture": True,
        }]}).encode())
        path = self.bundle / "text" / "D03.txt"
        path.write_bytes(path.read_bytes().rstrip(b"\n") + b"\n[... 8 characters not shown on the page ...]\n")
        gaps = cq.source_completeness(cq.Bundle(self.bundle))
        self.assertTrue(any("docs/c-notes.txt" in gap and "8 omitted characters" in gap for gap in gaps))
        self.assertIn("D03.txt: line 5, column 1, 8 omitted characters", gaps)
        self.assertEqual(self.codes(self.assessed_ready(self.rebound_ready(ready))), set())

    def test_record_parser_ignores_verdict_words_in_feedback(self):
        records = cq.records("- today | rev: 1 | platform quick test | platform_verdict: fail | "
                             "tiers: Low 5/5 · Medium — · High — | feedback: closed-book solve missed; outcome: match")
        self.assertEqual(len(records), 1)
        self.assertEqual(records[0]["kind"], "quick-test")
        self.assertEqual(records[0]["platform_verdict"], "fail")
        self.assertEqual(records[0]["reached"], ["easy"])
        injected = cq.records("- today | rev: 1 | platform quick test | platform_verdict: fail | "
                              "feedback: closed-book solve missed | platform_verdict: pass | outcome: match")
        self.assertEqual(injected[0]["platform_verdict"], "fail")
        self.assertIn("platform_verdict: pass", injected[0]["feedback"])
        forged = cq.records("- today | rev: 1 | tested_sha256: abc | answer: a quoted answer | outcome: match")
        self.assertEqual(forged, [])
        late_binding = cq.records("- today | rev: 1 | answer: a quoted answer | tested_sha256: abc | "
                                  "outcome: match | solver: A")
        self.assertEqual(late_binding, [])
        self.assertEqual(cq.records("- today | rev: 1oops | answer: unknown | outcome: match")[0]["rev"], None)

    def test_hook_reports_only_for_question_files(self):
        bad = self.write(GOOD.replace("documents: D01, D02", "documents: D01"))

        def run_hook(file_path: str) -> str:
            data = json.dumps({"tool_input": {"file_path": file_path}})
            with contextlib.redirect_stdout(io.StringIO()) as out, mock.patch("sys.stdin", io.StringIO(data)):
                cq.hook()
            return out.getvalue()

        note = json.loads(run_hook(str(bad)))["hookSpecificOutput"]["additionalContext"]
        self.assertIn("FAIL documents", note)
        self.assertEqual(run_hook(str(self.root / "tasks" / "template-task" / "questions" / "q00.md")), "")
        self.assertEqual(run_hook(str(self.bundle / "notes" / "doc-map.md")), "")

    def test_template_is_skipped_and_parses(self):
        template = TOOLS.parent / "tasks" / "template-task" / "questions" / "q00.md"
        meta, sections, found = cq.check(template, cq.Bundle(self.bundle))
        self.assertEqual(meta["status"], "template")
        self.assertEqual(found, [])
        self.assertEqual(sections["platform validation"], "")
        self.assertEqual(cq.platform_records(sections), [])
        self.assertNotIn("blind solve", sections)
        filled = template.read_bytes().decode("utf-8").replace("status: template", "status: draft")
        self.assertIn("fields", self.codes(filled))

    # split_capture.py
    def test_split_capture_writes_documents_only(self):
        dot = "·"
        capture = self.root / "capture.md"
        capture.write_bytes((
            f"# AQ-TEST {dot} Papers\n\n## Bundle overview\n\n### Doc 1 {dot} Overview title\n\n"
            "Summary text, not a document.\n\n## Documents\n\n"
            f"### Doc 1 {dot} paper_a.txt\n\n**Paper A**\n\n10 tokens {dot} 40 characters\n\n"
            "First paragraph of A.\n\n| a table-looking line\n\n"
            f"### Doc 2 {dot} paper_b.txt\n\n**Paper B**\n\n12 tokens {dot} 60 characters\n\n"
            "> **Incomplete:** only 50 of 60 characters were on the page.\n\nText of B.\n\n"
            "[… 10 characters not shown on the page …]\n\n"
            f"## Guide {dot} Question-Writing Guide\n\nRules, not a document.\n").encode("utf-8"))
        bundle = self.root / "tasks" / "papers"
        report = split_capture.run(capture, bundle)
        docs = bundle / "docs"
        b_text = "Text of B.\n\n[… 10 characters not shown on the page …]\n".encode("utf-8")
        self.assertEqual(sorted(p.name for p in docs.iterdir()), ["paper_a.txt", "paper_b.txt"])
        self.assertEqual((docs / "paper_a.txt").read_bytes(),
                         b"First paragraph of A.\n\n| a table-looking line\n")
        self.assertEqual((docs / "paper_b.txt").read_bytes(), b_text)
        self.assertIn("INCOMPLETE in the capture: only 50 of 60", report[1])
        self.assertIn("1 gap marker (10 characters, all of the shortfall) at line 3", report[1])
        self.assertNotIn("gap marker", report[0])
        self.assertEqual(split_capture.gap_lines({"body": "x\n", "note": "only 5 of 9 characters"}),
                         ["      no gap markers: the missing characters are not marked in the text"])
        self.assertIn("2 gap markers (7 characters, the note says 9 short) at lines 1, 3", split_capture.gap_lines(
            {"body": "[... 3 characters not shown on the page ...]\n\n[… 4 characters not shown on the page …]",
             "note": "only 1 of 10 characters"})[0])
        self.assertTrue(all("[unchanged]" in line.split("\n")[0] for line in split_capture.run(capture, bundle)))
        (docs / "paper_b.txt").write_bytes(b"edited")
        self.assertIn("left alone", split_capture.run(capture, bundle)[1])
        self.assertEqual((docs / "paper_b.txt").read_bytes(), b"edited")
        split_capture.run(capture, bundle, force=True)
        self.assertEqual((docs / "paper_b.txt").read_bytes(), b_text)
        unsafe = capture.read_bytes().replace(b"paper_a.txt", b"../escape.txt")
        with self.assertRaises(ValueError):
            split_capture.parse_capture(unsafe.decode("utf-8"))
        with self.assertRaises(ValueError):
            split_capture.parse_capture("# Not a capture\n")
        dated = split_capture.parse_capture(
            f"## Documents\n\n### Doc 1 {dot} 2024-04-26 {dot} cl_ex991_2024-04-26.txt\n\n"
            f"**Release**\n\n5 tokens {dot} 20 characters\n\nRelease text.\n")
        self.assertEqual([(d["name"], d["title"]) for d in dated], [("cl_ex991_2024-04-26.txt", "Release")])
        quoted = split_capture.parse_capture(
            f"## Documents\n\n### Doc 6 {dot} 2026-05-05 {dot} eog_ex991_2026-05-05.txt\n\n"
            f"**Release**\n\n17,678 tokens {dot} 53,814 characters {dot} 2/2 quotes found\n\nEX-99.1\n")
        self.assertEqual([(d["stated_chars"], d["body"]) for d in quoted], [(53814, "EX-99.1\n")])
        filed = split_capture.parse_capture(
            f"## Documents\n\n### Doc 1 {dot} 8-K {dot} 2026-04-23 {dot} fcx-1q-2026-ex-99-1.txt\n\n"
            f"**Release**\n\n25,168 tokens {dot} 99,686 characters\n\nFreeport Reports\n\n"
            f"### Doc 2 {dot} 6-K {dot} 2026-04-22 {dot} bhp-review.txt\n\n**Review**\n\n"
            f"15,475 tokens {dot} 42,162 characters\n\nUNITED STATES\n")
        self.assertEqual([(d["name"], d["title"], d["body"]) for d in filed],
                         [("fcx-1q-2026-ex-99-1.txt", "Release", "Freeport Reports\n"),
                          ("bhp-review.txt", "Review", "UNITED STATES\n")])
        undated = split_capture.parse_capture(
            f"## Documents\n\n### Doc 4 {dot} 8-K {dot} mchp-3q-2026-ex-99-1-file.txt\n\n"
            f"**Release**\n\n9,071 tokens {dot} 35,177 characters\n\nEXHIBIT 99.1\n")
        self.assertEqual([d["name"] for d in undated], ["mchp-3q-2026-ex-99-1-file.txt"])

    def test_split_capture_reads_plain_text_capture(self):
        dot, eq, rule = "·", "=" * 72, "-" * 72
        a_body = (f"Exhibit 99.1\n\n# # #\n\nPage one.\n-----\nPage two.\n{rule}\nNot a header line.\n"
                  f"{rule}\nDoc 9 {dot} 2026 results\nstill page two\n\n\n-----\n")
        capture = self.root / "capture.txt"
        capture.write_bytes((
            f"AQ-TEST {dot} Corporate filings\n\n{eq}\nBUNDLE OVERVIEW\n{eq}\n\n"
            f"Doc 1 {dot} Overview title\nSummary text, not a document.\n\n{eq}\nDOCUMENTS\n{eq}\n\n\n"
            f"{rule}\nDoc 1 {dot} 8-K {dot} 2026-02-18 {dot} adi-1q-2026.txt\nRelease One\n"
            f"10 tokens {dot} 40 characters\n{rule}\n\n{a_body}\n"
            f"{rule}\nDoc 2 {dot} 8-K {dot} mchp-3q-2026-file.txt\nRelease Two\n"
            f"12 tokens {dot} 60 characters\n{rule}\n\nText of two.\n\n\n-----\n\n"
            f"{rule}\nDoc 3 {dot} 2026-04-15 {dot} pgr_ex991_2026-04-15.txt\nRelease Three\n"
            f"5 tokens {dot} 20 characters {dot} 1/1 quotes found\n{rule}\n\nText of three.\n\n"
            f"{rule}\nDoc 4 {dot} 8-K {dot} fy2025 {dot} cor-q2-fy2025.txt\nRelease Four\n"
            f"6 tokens {dot} 24 characters\n{rule}\n\nText of four.\n\n"
            f"{eq}\nGUIDE\n{eq}\n\nRules, not a document.\n").encode("utf-8"))
        bundle = self.root / "tasks" / "filings"
        report = split_capture.run(capture, bundle)
        docs = bundle / "docs"
        self.assertEqual(sorted(p.name for p in docs.iterdir()),
                         ["adi-1q-2026.txt", "cor-q2-fy2025.txt", "mchp-3q-2026-file.txt",
                          "pgr_ex991_2026-04-15.txt"])
        self.assertEqual((docs / "adi-1q-2026.txt").read_bytes(), a_body.encode("utf-8"))
        self.assertEqual((docs / "mchp-3q-2026-file.txt").read_bytes(), b"Text of two.\n\n\n-----\n")
        self.assertEqual((docs / "pgr_ex991_2026-04-15.txt").read_bytes(), b"Text of three.\n")
        self.assertEqual((docs / "cor-q2-fy2025.txt").read_bytes(), b"Text of four.\n")
        self.assertEqual([line.split("  ")[2] for line in report],
                         ["40 chars", "60 chars", "20 chars", "24 chars"])
        parsed = split_capture.parse_capture(capture.read_bytes().decode("utf-8"))
        self.assertEqual([(d["number"], d["title"], d["note"]) for d in parsed],
                         [(1, "Release One", None), (2, "Release Two", None), (3, "Release Three", None),
                          (4, "Release Four", None)])
        intake = json.loads((bundle / "notes" / "intake.json").read_bytes())
        self.assertTrue(all(doc["matches_capture"] for doc in intake["documents"]))
        with self.assertRaises(ValueError):
            split_capture.parse_capture(f"{eq}\nBUNDLE OVERVIEW\n{eq}\n\nDoc 1 {dot} Overview\n")
        swapped = capture.read_bytes().decode("utf-8").replace(f"Doc 3 {dot} 2026", f"Doc 1 {dot} 2026")
        with self.assertRaises(ValueError):
            split_capture.parse_capture(swapped)

    # agent_scope.py
    def test_blind_solver_reads_only_documents(self):
        (self.root / ".claude").mkdir()
        (self.root / ".claude" / "solver-scope.json").write_bytes(
            json.dumps({"bundle": "freight", "run": "test-run"}).encode())
        def decide(tool: str, **tool_input) -> str | None:
            return agent_scope.check("blind-solver", {"tool_name": tool, "tool_input": tool_input,
                                                      "cwd": str(self.root)}, self.root)
        t = self.root / "tasks"
        packet = t / "freight" / "runs" / "test-run" / "solver-1"
        self.assertIsNone(decide("Read", file_path=str(packet / "text/D01.txt")))
        self.assertIsNone(decide("Read", file_path="tasks/freight/runs/test-run/solver-1/docs/a-release-q1.pdf"))
        self.assertIsNone(decide("Grep", pattern="revenue", path=str(packet / "text")))
        self.assertIsNone(decide("Glob", pattern="*.txt", path=str(packet / "text")))
        for denied in [("Read", {"file_path": str(t / "freight/questions/q01.md")}),
                       ("Read", {"file_path": str(t / "freight/text/D01.txt")}),
                       ("Read", {"file_path": str(t / "freight/docs/a-release-q1.pdf")}),
                       ("Read", {"file_path": str(t / "freight/notes/doc-map.md")}),
                       ("Read", {"file_path": str(t / "freight/text/../questions/q01.md")}),
                       ("Read", {"file_path": str(self.root / "dist/freight.md")}),
                       ("Read", {"file_path": str(t / "template-task/text/D01.txt")}),
                       ("Grep", {"pattern": "revenue"}),
                       ("Glob", {"pattern": "../questions/*", "path": str(t / "freight/text")}),
                       ("Edit", {"file_path": str(t / "freight/text/D01.txt")})]:
            with self.subTest(denied=denied):
                self.assertIsNotNone(decide(denied[0], **denied[1]))
        self.assertIsNotNone(agent_scope.check("unknown-role", {"tool_name": "Read"}, self.root))

    def test_scope_hook_end_to_end(self):
        def run(payload: dict) -> str:
            return subprocess.run([sys.executable, str(TOOLS / "agent_scope.py"), "blind-solver"],
                                  input=json.dumps(payload).encode(), capture_output=True, check=True).stdout.decode()
        root = TOOLS.parent
        denied = run({"tool_name": "Read", "tool_input": {"file_path": str(root / "tasks/x/questions/q01.md")}})
        self.assertEqual(json.loads(denied)["hookSpecificOutput"]["permissionDecision"], "deny")
        # A notes file is denied even if another test has prepared an active workspace scope.
        notes = json.loads(run({"tool_name": "Read", "tool_input": {"file_path": str(root / "tasks/x/notes/doc-map.md")}}))
        self.assertEqual(notes["hookSpecificOutput"]["permissionDecision"], "deny")
        broken = subprocess.run([sys.executable, str(TOOLS / "agent_scope.py"), "blind-solver"],
                                input=b"not json", capture_output=True, check=True).stdout.decode()
        self.assertIn("failed closed", broken)

    # The form's reasoning number check in the floor and the hand-over (tools/form_numbers.py).
    def test_form_number_misses_fail_the_floor_and_block_the_handover(self):
        self.write_platform_numbers()
        wrong_doc = GOOD.replace("Step 2. Source fact (Doc 1, revenue table)", "Step 2. Source fact (Doc 2, revenue table)")
        _, _, found = cq.check(self.write(wrong_doc), cq.Bundle(self.bundle))
        misses = [message for level, code, message in found if level == "FAIL" and code == "form-numbers"]
        self.assertEqual(len(misses), 2)
        self.assertTrue(misses[0].startswith("Step 2: Doc 2 doesn't mention 409.8. Doc 1 prints it"))
        with contextlib.redirect_stdout(io.StringIO()):
            self.assertEqual(cq.handover(self.bundle, ["q01"]), 1)
        sheet = (self.bundle / "notes" / "handover-q01.md").read_bytes().decode("utf-8")
        self.assertIn("NOT READY TO PASTE", sheet)
        self.assertIn("FAIL form-numbers: Step 2: Doc 2 doesn't mention 409.8.", sheet)
        # A ready question was measured with this reasoning: the miss is reported, not a reason for draft.
        ready = self.rebound_ready(self.targeted_ready(local=False).replace(
            "(Doc 1, revenue table)", "(Doc 2, revenue table)"))
        _, _, found = cq.check(self.write(ready), cq.Bundle(self.bundle))
        self.assertEqual([code for level, code, _ in found if level == "FAIL"], [])
        self.assertTrue(any(level == "WARN" and code == "form-numbers" and "new difficulty result" in message
                            for level, code, message in found))
        # A step that leans on an unticked document is reported (reviewer returns, aq-01333 q01 and aq-02749 q02).
        leaning = GOOD.replace("so revenue missed the forecast.",
                               "so revenue missed the forecast. Doc 3 (not needed for the count) says the same.")
        _, _, found = cq.check(self.write(leaning), cq.Bundle(self.bundle))
        self.assertTrue(any(level == "WARN" and code == "documents" and message.startswith(
            "the reasoning names Doc 3, which is not ticked as necessary") for level, code, message in found))
        _, _, found = cq.check(self.write(GOOD), cq.Bundle(self.bundle))
        self.assertFalse(any("is not ticked as necessary" in message for _, _, message in found))
        # Frozen questions are history, and without page numbers the check says it could not run.
        self.assertNotIn("form-numbers", self.codes(wrong_doc.replace("status: draft", "status: submitted"), "WARN"))
        (self.bundle / "notes" / "intake.json").unlink()
        self.assertIn("form-numbers", self.codes(GOOD, "WARN"))
        self.assertNotIn("form-numbers", self.codes(wrong_doc))


class FormNumbers(unittest.TestCase):
    """The form's reasoning number check against the form's recorded results (user screenshots)."""

    FILLER = "Background paragraph on enforcement of the import ban. " * 700  # about 38,500 characters

    def test_totals_printed_only_in_words_are_flagged_as_the_form_did(self):
        # aq-02505 q01 rev 0 (2026-10-09): "Step 1: Doc 2 doesn't mention 107." and "Step 7: Doc 6 doesn't
        # mention 187." The notices print both totals only in words; "five" early in Doc 2 passed as 5.
        doc2 = ("Federal Register (Monday, November 25, 2024). See appendix 1. This update adds five entities to "
                "one list and twenty-five entities to another. One of them was first listed in June 2022. "
                + self.FILLER + "One hundred and seven entities that meet the criteria set forth in the four "
                "required lists")
        doc6 = ("The list as of August 3, 2026 is in appendix 1. " + self.FILLER + "One hundred and eighty-seven "
                "entities that meet the criteria set forth in the four required lists")
        reasoning = "\n\n".join([
            'Step 1. Source fact (Doc 2, appendix 1): the list as of November 25, 2024 specifies "One hundred and '
            'seven entities" (107).',
            "Step 2. Source fact (Doc 2, summary): the update adds five entities (5) to one list and twenty-five "
            "entities (25) to another.",
            "Step 3. Source fact (Doc 2): one of the twenty-five was first listed in June 2022.",
            "Step 4. Calculation (Doc 2 and Doc 6): additions: 5 + 25 = 30.",
            "Step 5. Calculation (Doc 2 and Doc 6): entities new to the list: 30 - 1 = 29.",
            "Step 6. Calculation (Doc 2 and Doc 6): entities on the earlier version: 107 - 29 = 78.",
            'Step 7. Source fact (Doc 6, appendix 1): the list as of August 3, 2026 specifies "One hundred and '
            'eighty-seven entities" (187).',
            "Step 8. Calculation (Doc 2 and Doc 6): growth: 187 - 78 = 109.",
            "Step 9. Calculation (Doc 2 and Doc 6): growth as a percentage of the earlier count: 109 / 78 = 139.7%."])
        found = fn.check(reasoning, {2: doc2, 6: doc6})
        self.assertEqual([(f.level, f.form) for f in found], [("FAIL", "Step 1: Doc 2 doesn't mention 107."),
                                                              ("FAIL", "Step 7: Doc 6 doesn't mention 187.")])
        self.assertIn('only in words ("One hundred and seven"', found[0].why)
        self.assertIn("quote it as evidence", found[0].why)
        # Quoted as evidence, the totals are the documents' own statements: expected false alarms, which
        # the user knew were wrong warnings ("this time, it is not your fault", 2026-10-09).
        quotes = {2: ["One hundred and seven entities that meet the criteria"],
                  6: ["One hundred and eighty-seven entities that meet the criteria"]}
        found = fn.check(reasoning, {2: doc2, 6: doc6}, quotes)
        self.assertEqual([(f.level, f.form) for f in found], [("WARN", "Step 1: Doc 2 doesn't mention 107."),
                                                              ("WARN", "Step 7: Doc 6 doesn't mention 187.")])
        self.assertTrue(all(f.why.startswith("expected false alarm") for f in found))
        # A quote from the other document does not excuse the citation.
        self.assertEqual([f.level for f in fn.check(reasoning, {2: doc2, 6: doc6}, {6: quotes[2]})], ["FAIL", "FAIL"])

    def test_recorded_rules(self):
        doc1 = ("Table 2: 1,295 individuals of 1,301 responses. Revenue is expected at $3.15 billion and gross "
                "margin at 60.7%. Page 2 of 4.")
        doc4 = "Paragraph 3.9 and Table 2: 102 individuals answered and 10 more did not. Annex 1."
        docs = {1: doc1, 4: doc4}

        def forms(reasoning: str, level: str = "FAIL") -> list[str]:
            return [f.form for f in fn.check(reasoning, docs) if f.level == level]

        derived = "Step 1. Calculation (Doc 4): individuals: 102 + 10 = 112.\n\n"
        # aq-02102 q01 rev 1: a step citing one document accepts no earlier result; two documents do.
        self.assertEqual(forms(derived + "Step 2. Calculation (Doc 4): half: 112 / 2 = 56."),
                         ["Step 2: Doc 4 doesn't mention 112."])
        self.assertEqual(forms(derived + "Step 2. Calculation (Doc 1 and Doc 4): half: 112 / 2 = 56."), [])
        # aq-02486 q01 rev 0: a new number in a step citing no document.
        self.assertEqual(forms("Step 1. The request covers 225 pages."),
                         ["Step 1: Which document says 225? Add 'Doc N'."])
        # aq-01960 q02: unit words inside the expression leave the result unexempted; a label before the
        # bare expression and a longer chain passed.
        self.assertEqual(forms("Step 1. Calculation (Doc 1): 3.15 billion x 60.7% = 1.91205 billion."),
                         ["Step 1: Doc 1 doesn't mention 1.91205."])
        self.assertEqual(forms("Step 1. Calculation (Doc 1): gross profit in billions of dollars = 3.15 × 60.7% = "
                               "1.91205."), [])
        self.assertEqual(forms(derived + "Step 2. Calculation (Doc 1 and Doc 4): total = 112 + 1,295 + 2 = 1,409."), [])
        # aq-02626 q03 Step 4 ("twelve 30-day months ...: 12 × 30 + 5 = 365") and aq-01321 q02 Step 9 ("Twelve
        # months ended June 30, 2026: 1,942 + 1,999 = 3,941") were flagged on exact results; both steps spell a
        # number out. aq-02626 q03 Step 5, a mixed chain with no number word, drew no line: a WARN only.
        wordy = fn.check("Step 1. Calculation (Doc 1 and Doc 4): the twelve responses: 102 + 10 = 112.", docs)
        self.assertEqual([(f.level, f.form) for f in wordy], [("FAIL", "Step 1: Docs 1 and 4 don't mention 112.")])
        self.assertIn("spells a number out", wordy[0].why)
        mixed = fn.check("Step 1. Calculation (Doc 1 and Doc 4): 102 × 2 + 10 = 214.", docs)
        self.assertEqual([(f.level, f.form) for f in mixed], [("WARN", "Step 1: a chain that mixes operators.")])
        self.assertEqual(forms("Step 1. Calculation (Doc 1 and Doc 4): 102 × 2 = 204.\n\n"
                               "Step 2. Calculation (Doc 1 and Doc 4): 204 + 10 = 214."), [])
        # aq-02157 q01: a result typed as a percentage is not reusable as a plain number.
        typed = forms("Step 1. Calculation (Doc 1 and Doc 4): share: 1,295 / 1,301 = 99.54%.\n\n"
                      "Step 2. Calculation (Doc 1 and Doc 4): gap: 99.54 - 2 = 97.54.")
        self.assertEqual(typed, ["Step 2: Docs 1 and 4 don't mention 99.54."])
        # aq-00949 q04 / rev 3 of aq-00745: a restated or rounded result is read as a document figure.
        restated = fn.check("Step 1. Calculation (Doc 1): share: 1,295 / 1,301 = 99.539%, or 99.5%.", docs)
        self.assertEqual([f.number for f in restated], ["99.5%"])
        self.assertIn("restates", restated[0].why)
        self.assertEqual(forms("Step 1. Calculation (Doc 1): 1,295 / 1,301 = 0.9954 (rounded)."),
                         ["Step 1: Doc 1 doesn't mention 0.9954."])
        # A figure from another document: the diagnosis names the document that prints it.
        wrong = fn.check("Step 1. Source fact (Doc 1): 102 individuals answered.", docs)
        self.assertEqual([f.form for f in wrong], ["Step 1: Doc 1 doesn't mention 102."])
        self.assertIn("Doc 4 prints it", wrong[0].why)
        # A range reads as its first document (aq-01960 q02); step and document references are not figures.
        ranged = fn.check(derived + "Step 2. Calculation (from Steps 1 to 1; Docs 1 to 4): 112 + 2 = 114.", docs)
        self.assertEqual([(f.level, f.number) for f in ranged], [("WARN", ""), ("FAIL", "112")])

    def test_depth_limits(self):
        # aq-02157: printed digits near character 80,000 were flagged; 19,581 is the deepest observed pass.
        deep = {2: self.FILLER * 2 + "The minimum duty is 8.46 pounds.",
                3: self.FILLER[:25_000] + "Rates fell by 24 points."}
        found = fn.check("Step 1. Source fact (Doc 2): 8.46 pounds.\n\nStep 2. Source fact (Doc 3): 24 points.", deep)
        self.assertEqual([(f.level, f.form) for f in found], [("FAIL", "Step 1: Doc 2 doesn't mention 8.46."),
                                                              ("WARN", "Step 2: Doc 3 doesn't mention 24.")])
        self.assertIn("past the 60,000 characters", found[0].why)
        quoted = fn.check("Step 1. Source fact (Doc 2): 8.46 pounds.", deep, {2: ["The minimum duty is 8.46 pounds."]})
        self.assertEqual([(f.level, f.why.split(":")[0]) for f in quoted], [("WARN", "expected false alarm")])

    def test_sketch_input_notes(self):
        text = "Summary. " + self.FILLER + "One hundred and seven entities are listed; 8.46 pounds per pack."
        notes = fn.input_notes("One hundred and seven entities are listed; 8.46 pounds per pack", text)
        self.assertEqual(len(notes), 2)
        self.assertIn("states 107 only in words", notes[0])
        self.assertIn("8.46 is printed first at character", notes[1])
        self.assertEqual(fn.input_notes("adds five entities to one list", "This update adds five entities to one list."), [])


if __name__ == "__main__":
    unittest.main(verbosity=1)
